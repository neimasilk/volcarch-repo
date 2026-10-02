"""P23 human validation: fixed random sample of village-noun occurrences, blank sheet for an epigrapher.

The sheet shows only the context (no AI codes, no proposed names), so the human coding is blind.
Metrics are fixed in PROTOCOL.md before any human code exists.
"""
import csv
import os
import random

import openpyxl
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "..", "..", "experiments", "E226_t7_detection_calibration", "frame", "coder_input.csv")
rows = list(csv.DictReader(open(SRC, encoding="utf-8")))
rng = random.Random(23)
sample = sorted(rng.sample(rows, 100), key=lambda r: r["occ_id"])
with open(os.path.join(HERE, "sample_ids.csv"), "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh)
    w.writerow(["occ_id"])
    w.writerows([[r["occ_id"]] for r in sample])
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "lembar_kode"
ws.append(["no", "id", "konteks kiri", "KATA", "konteks kanan",
           "menyebut desa tertentu? (Y/N/?)", "nama desa (tanpa partikel i/ri)", "catatan"])
for i, r in enumerate(sample, 1):
    ws.append([i, r["occ_id"], r["kwic_left"], r["noun"], r["kwic_right"], "", "", ""])
for c in ws[1]:
    c.font = Font(bold=True)
for col, wd in zip("ABCDEFGH", [5, 26, 60, 10, 60, 14, 26, 30]):
    ws.column_dimensions[col].width = wd
for row in ws.iter_rows(min_row=2):
    for c in row[2:5]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
g = wb.create_sheet("petunjuk")
for line in [
    "Setiap baris adalah satu kemunculan kata wanua/banua/wanwa/thāni dalam teks edisi DHARMA (prasasti 701–1000 M).",
    "Kolom F: Y bila kata itu diikuti nama desa tertentu (mis. 'anak wanua i X'); N bila umum/tanpa nama (mis. 'anak wanua kabaiḥ', gelar seperti 'rake wanua poḥ'); ? bila ragu.",
    "Kolom G: tulis nama desa sebagaimana di teks (huruf kecil, tanpa partikel i/ri/riṅ), berhenti sebelum 'watak/vatak' atau gelar/orang berikutnya.",
    "Ejaan DHARMA: A/I/U kapital = vokal awal; v = w; ṅ = ng. Tanda '-' di tengah kata = pemenggalan baris.",
    "Mohon jangan melihat daftar desa yang dilampirkan sebelumnya saat mengisi lembar ini, supaya penilaian Ibu/Bapak independen.",
]:
    g.append([line])
g.column_dimensions["A"].width = 150
wb.save(os.path.join(HERE, "lembar_validasi_100_kemunculan.xlsx"))
print(len(sample), "rows; first:", sample[0]["occ_id"])
