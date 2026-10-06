# -*- coding: utf-8 -*-
"""Build one Word file with the revision tables A-H; every number is written from the CSV files.

Why this exists: the PI pastes real Word tables into the journal template. Nothing here is typed by hand:
cells are produced by fmt() from the CSV value, and after saving the .docx is re-opened and every numeric cell
is compared with the CSV value rounded by the same rule (readback_check.json). Exit code 1 if any cell differs.
Rounding is half-up on the decimal repr of the CSV value (so 0.7103 -> 0.710, 0.0005 -> 0.001).
No manuscript text, no interpretation: captions are placeholders.
"""
import csv, json, re, sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = Path(__file__).resolve().parent
EXP = HERE.parent
REPO = EXP.parent
OUT = HERE / "results" / "tables_docx"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "P8_revision_tables.docx"
DAG = "†"
DDAG = "‡"


def rel(p):
    return str(p.relative_to(REPO)).replace("\\", "/")


def readcsv(p):
    with open(p, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


INP = {"FM25": "form + meaning (25)", "F17": "form only (17)"}


# --- rounding / formatting: the only place numbers become text ---------------------------------------
def dec(v):
    return Decimal(str(v).strip())


def q(v, nd):
    return dec(v).quantize(Decimal(1).scaleb(-nd), ROUND_HALF_UP)


def fmt(v, rule):
    if rule == "int":
        return str(int(round(float(v))))
    if rule == "d3":
        return f"{q(v, 3)}"
    if rule == "d2":
        return f"{q(v, 2)}"
    if rule == "pct1":
        return f"{q(v, 1)}"
    if rule == "frac_pct1%":
        return f"{(dec(v) * 100).quantize(Decimal('0.1'), ROUND_HALF_UP)}%"
    if rule in ("sd3", "sd2"):
        x = q(v, int(rule[2]))
        # a value that rounds to zero is printed unsigned: it has no sign
        if x == 0:
            x = abs(x)
        return ("+" if x > 0 else "") + f"{x}"
    raise ValueError(rule)


def expected(v, rule):
    """Decimal the cell must show, computed from the raw CSV string, independent of the cell text."""
    if rule == "int":
        return Decimal(int(round(float(v))))
    if rule in ("d3", "sd3"):
        return q(v, 3)
    if rule in ("d2", "sd2"):
        return q(v, 2)
    if rule == "pct1":
        return q(v, 1)
    if rule == "frac_pct1%":
        return (dec(v) * 100).quantize(Decimal("0.1"), ROUND_HALF_UP)
    raise ValueError(rule)


class Num:
    def __init__(self, v, rule):
        self.v, self.rule = v, rule


def N(v, rule):
    return Num(v, rule)


def cell_text(c):
    if isinstance(c, str):
        return c
    if isinstance(c, Num):
        return fmt(c.v, c.rule)
    return "".join(p if isinstance(p, str) else fmt(p.v, p.rule) for p in c)


def cell_nums(c):
    if isinstance(c, str):
        return []
    if isinstance(c, Num):
        return [c]
    return [p for p in c if isinstance(p, Num)]


T = []

# A
p = HERE / "results/T1_label_by_list.csv"
R = readcsv(p)
T.append(dict(id="A", name="label by list", src=[p],
              header=["List", "Forms", "Coded (n)", "Coded (%)", "Candidate (n)", "Candidate (%)"],
              rows=[[r["list"], N(r["forms"], "int"), N(r["coded_n"], "int"), N(r["coded_pct"], "pct1"),
                     N(r["candidate_n"], "int"), N(r["candidate_pct"], "pct1")] for r in R], notes=[]))
# B
p = HERE / "results/T2_cv_performance.csv"
R = readcsv(p)
T.append(dict(id="B", name="cross-validated performance", src=[p],
              header=["Inputs", "Classifier", "AUC", "SD of AUC, 10 seed means", "SD of AUC, 50 folds", "Accuracy",
                      "Always-coded accuracy", "Precision (candidate)", "Recall (candidate)", "F1 (candidate)"],
              rows=[[INP[r["input_set"]], r["classifier"], N(r["auc_mean"], "d3"), N(r["auc_sd_seed_means"], "d3"),
                     N(r["auc_sd_folds"], "d3"), N(r["accuracy_mean"], "d3"), N(r["accuracy_always_coded"], "d3"),
                     N(r["precision_candidate_mean"], "d3"), N(r["recall_candidate_mean"], "d3"),
                     N(r["f1_candidate_mean"], "d3")] for r in R], notes=[]))
# C
p = HERE / "results/T3_lolo.csv"
R = readcsv(p)
fm = {r["list_key"]: r for r in R if r["input_set"] == "FM25" and r["list_key"] != "SUMMARY"}
f17 = {r["list_key"]: r for r in R if r["input_set"] == "F17" and r["list_key"] != "SUMMARY"}
sm = {r["input_set"]: r for r in R if r["list_key"] == "SUMMARY"}
rows = []
for k, r in fm.items():
    s = f17[k]
    assert r["accuracy_majority_answer"] == s["accuracy_majority_answer"] and r["n"] == s["n"]
    rows.append([r["list"], N(r["n"], "int"), N(r["n_candidates"], "int"), N(r["auc"], "d3"), N(r["accuracy"], "d3"),
                 N(s["auc"], "d3"), N(s["accuracy"], "d3"), N(r["accuracy_majority_answer"], "d3")])
rows.append(["Mean (SD) over the six lists " + DAG, "", "",
             [N(sm["FM25"]["auc"], "d3"), " (", N(sm["FM25"]["auc_sd_over_lists"], "d3"), ")"], "",
             [N(sm["F17"]["auc"], "d3"), " (", N(sm["F17"]["auc_sd_over_lists"], "d3"), ")"], "", ""])
T.append(dict(id="C", name="one list held out", src=[p],
              header=["Held-out list", "n", "Candidates", "AUC, form + meaning (25)", "Accuracy, form + meaning (25)",
                      "AUC, form only (17)", "Accuracy, form only (17)", "Majority-answer accuracy"],
              rows=rows, notes=[DAG + " " + sm["FM25"]["list"] + " (CSV row SUMMARY)."]))
# D
p = HERE / "results/T5_cells_by_list.csv"
R = readcsv(p)
T.append(dict(id="D", name="label by profile cells per list", src=[p],
              header=["List", "n", "Candidate and profile", "Candidate and no profile", "Coded and profile",
                      "Coded and no profile", "Kappa"],
              rows=[[r["list"], N(r["n"], "int"), N(r["candidate_and_profile"], "int"),
                     N(r["candidate_and_no_profile"], "int"), N(r["coded_and_profile"], "int"),
                     N(r["coded_and_no_profile"], "int"), N(r["kappa"], "d3")] for r in R], notes=[]))
# E
p = HERE / "results/T6_profile.csv"
R = readcsv(p)
rows = []
types = {}
for r in R:
    binary = r["type"] == "0/1"
    mean_rule = "frac_pct1%" if binary else "d2"
    is_or = r["effect_type"].startswith("Mantel")
    tag = DDAG if is_or else DAG
    types[tag] = r["effect_type"]
    er = "d2" if is_or else "sd2"
    rows.append([r["property"], N(r["mean_candidate"], mean_rule), N(r["mean_coded"], mean_rule),
                 N(r["n_lists_same_sign_as_pooled"], "int"),
                 ("Odds ratio " if is_or else "Mean difference ") + tag,
                 N(r["effect"], er), N(r["ci_lo"], er), N(r["ci_hi"], er)])
T.append(dict(id="E", name="profile of candidates against coded forms", src=[p],
              header=["Property", "Candidates", "Coded forms", "Lists with the same sign as pooled (of 6)",
                      "Effect type", "Effect", "95% CI, lower", "95% CI, upper"],
              rows=rows, notes=[f"{t} {types[t]}." for t in sorted(types)]))
# F
p = EXP / "E228_p8_revision_analyses/results/TABLE_R1-2_retention_by_language.csv"
R = readcsv(p)
T.append(dict(id="F", name="retention from Proto-Malayo-Polynesian by list", src=[p],
              header=["List", "Meanings compared with PMP", "Retained from PMP (n)", "Retained from PMP (%)",
                      "Coded, not a PMP etymon (n)", "Coded, not a PMP etymon (%)", "No cognate set (n)",
                      "No cognate set (%)", "Not retained (%)", "Retained counting doubtful (%)",
                      "Meanings compared with PAn", "Retained from PAn (%)"],
              rows=[[r["language"], N(r["meanings_compared_with_PMP"], "int"), N(r["retained_from_PMP_n"], "int"),
                     N(r["retained_from_PMP_pct"], "pct1"), N(r["coded_not_PMP_etymon_n"], "int"),
                     N(r["coded_not_PMP_etymon_pct"], "pct1"), N(r["no_cognate_set_n"], "int"),
                     N(r["no_cognate_set_pct"], "pct1"), N(r["not_retained_pct"], "pct1"),
                     N(r["retained_pct_counting_doubtful"], "pct1"), N(r["meanings_compared_with_PAn"], "int"),
                     N(r["retained_from_PAn_pct"], "pct1")] for r in R], notes=[]))
# G (convention names as listed in the E228 README, section S1)
CONV = {"V0_as_published": "as published", "V1_q": "every marker to q", "V2_k": "every marker to k",
        "V3_unwritten": "marker not written", "V4_geminate": "geminate before consonant, else not written",
        "V5_all_as_glottal_letter": "apostrophe to glottal-stop letter", "V6_feature_removed": "glottal input removed"}
p = EXP / "E228_p8_revision_analyses/results/TABLE_R1-9_glottal_conventions.csv"
R = readcsv(p)
T.append(dict(id="G", name="glottal-stop conventions", src=[p],
              header=["Convention", "Forms changed", "CV AUC, 25 inputs", "Difference, 25 inputs",
                      "Leave-one-list-out mean AUC, 25 inputs", "Difference, 25 inputs (leave-one-list-out)",
                      "CV AUC, 26 inputs", "Difference, 26 inputs", "Leave-one-list-out mean AUC, 26 inputs"],
              rows=[[r["convention"] + " (" + CONV[r["convention"]] + ")", N(r["forms_changed"], "int"),
                     N(r["cv_auc_25"], "d3"), N(r["delta_cv_25"], "sd3"), N(r["lolo_mean_25"], "d3"),
                     N(r["delta_lolo_25"], "sd3"), N(r["cv_auc_26"], "d3"), N(r["delta_cv_26"], "sd3"),
                     N(r["lolo_mean_26"], "d3")] for r in R], notes=[]))
# H (variant descriptions as listed in the E230 README)
VAR = {"V0": "baseline", "D1": "digraphs to single symbols",
       "D2": "digraphs to single symbols, initial-letter inputs kept",
       "L1": "length = number of vowel groups", "L2": "form_length removed",
       "L3": "form_length and n_vowels removed"}
p = EXP / "E230_p8_revisit_E041_E042_digraph_length/results/B_variants.csv"
R = readcsv(p)
T.append(dict(id="H", name="digraph and length variants", src=[p],
              header=["Variant", "Inputs", "CV AUC", "Difference from the baseline", "Leave-one-list-out mean AUC",
                      "Difference for a held-out list"],
              rows=[[r["variant"] + " (" + VAR[r["variant"]] + ")", INP[r["model"]], N(r["cv_auc"], "d3"),
                     N(r["cv_delta_vs_V0"], "sd3"), N(r["lolo_mean"], "d3"), N(r["lolo_delta_vs_V0"], "sd3")]
                    for r in R], notes=[]))

# --- Word writing ------------------------------------------------------------------------------------
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for a in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, a, Cm(1.8))
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(8)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(2)


def para(text, bold=False, size=8, after=2):
    p_ = doc.add_paragraph()
    r_ = p_.add_run(text)
    r_.bold = bold
    r_.font.size = Pt(size)
    r_.font.name = "Times New Roman"
    p_.paragraph_format.space_after = Pt(after)
    return p_


def border(cell, edge):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), "6")
    e.set(qn("w:color"), "000000")
    b.append(e)


para("Revision tables A-H for P8 (Oceanic Linguistics), as Word tables for pasting into the journal template.", True)
para("Generated by experiments/E229_p8_revision_tables/06_tables_docx.py on 2026-10-06 from the CSV files named under each table.")
para("This file contains no manuscript text; the captions are placeholders.")

for t in T:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(f"TABLE {t['id']}. [CAPTION TO BE WRITTEN BY THE AUTHORS]".upper(), True, 8, 3)
    tb = doc.add_table(rows=1 + len(t["rows"]), cols=len(t["header"]))
    for j, h in enumerate(t["header"]):
        c = tb.rows[0].cells[j]
        c.text = ""
        r_ = c.paragraphs[0].add_run(h)
        r_.bold = True
        r_.font.size = Pt(8)
        r_.font.name = "Times New Roman"
        border(c, "bottom")
    for ri, row in enumerate(t["rows"], 1):
        for j, cv in enumerate(row):
            c = tb.rows[ri].cells[j]
            c.text = ""
            r_ = c.paragraphs[0].add_run(cell_text(cv))
            r_.font.size = Pt(8)
            r_.font.name = "Times New Roman"
            if ri == len(t["rows"]):
                border(c, "bottom")
    for rw in tb.rows:
        for c in rw.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
    for n in t["notes"]:
        para(n, False, 8, 1)
    para("Source: " + ", ".join(rel(s) for s in t["src"]), False, 7, 0)
doc.save(DOCX)

# --- read-back: re-open the docx, compare every numeric cell with the CSV value rounded by the rule ----
d2 = Document(DOCX)
tabs = d2.tables
assert len(tabs) == len(T), (len(tabs), len(T))
NUM = re.compile(r"[+-]?\d+(?:\.\d+)?")
report = {}
bad = 0
for t, tb in zip(T, tabs):
    n = m = 0
    details = []
    assert len(tb.rows) == 1 + len(t["rows"])
    for ri, row in enumerate(t["rows"], 1):
        for j, cv in enumerate(row):
            nums = cell_nums(cv)
            if not nums:
                continue
            shown = tb.rows[ri].cells[j].text
            found = [Decimal(x) for x in NUM.findall(shown)]
            for k, nm in enumerate(nums):
                n += 1
                exp = expected(nm.v, nm.rule)
                if not (k < len(found) and found[k] == exp and len(found) == len(nums)):
                    m += 1
                    details.append(dict(row=ri, col=t["header"][j], shown=shown, expected=str(exp)))
    report[t["id"]] = dict(table=t["name"], source_csv=[rel(s) for s in t["src"]], rows=len(t["rows"]),
                           cols=len(t["header"]), numeric_cells_checked=n, mismatches=m, details=details)
    bad += m

json.dump(dict(rounding="half-up on the decimal repr; AUC/accuracy/F1/precision/recall/kappa 3 dp; percentages 1 dp; "
                        "means, mean differences, odds ratios 2 dp; counts integer; differences signed "
                        "(a value that rounds to zero is unsigned)",
               total_mismatches=bad, tables=report),
          open(OUT / "readback_check.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for k, v in report.items():
    print(k, v["table"], v["rows"], "x", v["cols"], "numeric cells", v["numeric_cells_checked"], "mismatches",
          v["mismatches"])
print("TOTAL MISMATCHES", bad)
sys.exit(1 if bad else 0)
