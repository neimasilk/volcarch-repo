"""E226 step 1a: village-noun occurrences in 8th-10th c. DHARMA editions, with KWIC, for two-coder curation.

Consolidates the 2026-10-01 desk-scan scratch scripts (t7_parse/t7_main/t7_tab/t7_final) into one
reproducible script. Rules are frozen in ../DESIGN.md §2; change them there first, not here.
"""
import collections
import csv
import glob
import json
import os
import re
import sys
import unicodedata

import openpyxl
from lxml import etree

sys.stdout.reconfigure(encoding="utf-8")

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
DH = os.path.join(ROOT, "experiments", "E023_ritual_screening", "data", "dharma")
E082 = os.path.join(ROOT, "experiments", "E082_inscription_georeferencing", "results", "canonical30",
                    "geocoded_inscriptions_canonical30.csv")
OUT = os.path.join(os.path.dirname(__file__), "..", "frame")
TEI = "{http://www.tei-c.org/ns/1.0}"
DROP = {"sic", "orig", "del", "surplus", "note", "rdg", "fw", "label"}

# Undated-by-title items resolved by hand on 2026-10-01 against the IDENK xlsx; kept explicit so the
# sensitivity run (title-dated only) can drop them.
ALIAS = {
    "INSIDENKPoh": (905, "alias: xlsx Poh/Randusari I 827 Saka"),
    "INSIDENKKurungan": (885, "alias: xlsx Randusari II 807 Saka (?)"),
    "INSIDENKPanggumulanA": (902, "alias: xlsx Panggumulan I/A 824 Saka"),
    "INSIDENKPasanggrahan": ((900, 910), "alias: title ca. 900-910"),
    "INSIDENKWunganDaik": ((901, 1000), "alias: xlsx Wunandaik 10th c."),
    "INSIDENKJamwi": ((901, 1000), "alias: xlsx Sinaguha 10th c.? uncertain"),
    "INS12Gilikan": ((901, 1000), "alias: xlsx Gilikan I 10th c."),
    "INSIDENKHarinjing": ((804, 927), "alias: xlsx Harinjing A/B/C 804/921/927 ambiguous"),
    "INSIDENKWuruduKidul": (922, "alias: xlsx Wurudu Kidul A 844 Saka"),
}
# Large charters whose title does not match the xlsx designation; region looked up under this name.
XLSX_NAME = {
    "INSIDENKPoh": "Poḥ/Randusari I",
    "INSIDENKPanggumulanA": "Panggumulan I/A/Kembang Arum A",
    "INSIDENKKurungan": "Kurungan/Randusari II/Dang Acaryya Munindra",
    "INSIDENKWuruduKidul": "Wurudu Kidul A",
}
KEDU = ("magelang", "temanggung", "wonosobo", "purworejo")
PRAM = ("sleman", "bantul", "yogyakarta", "klaten", "gunung kidul", "gunungkidul", "kulon")

# b/v alternation: "banuA" is a vanua spelling (amendment A1, DESIGN §6, before coding).
VILLAGE = re.compile(r"^([vb]anua|[vb]anu|[vb]anun|[vb]anuan|[vb]anuana|[vb]anuanira|[vb]anuanya|[vb]anva|"
                     r"[vb]anvaa|[vb]anvanya|[vb]anvaniram|thani|thaninya)$")
PART = {"i", "ri", "rin", "ing", "in", "ni", "rim", "im"}
TOK = re.compile(r"[^\s.,;:!?()\[\]{}|/<>\"“”‘’·=–—…§‚]+")


def norm(s):
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.replace("w", "v")


def key(s):
    return re.sub(r"[^a-z0-9]", "", norm(str(s)))


def walk(e, out):
    for ch in e:
        if not isinstance(ch.tag, str):
            if ch.tail:
                out.append(ch.tail)
            continue
        ln = etree.QName(ch).localname
        if ln in DROP:
            pass
        elif ln == "lb" and ch.get("break") == "no":
            if ch.tail:
                out.append(ch.tail.lstrip())
            continue
        elif ln in ("lb", "pb", "milestone"):
            out.append(" ")
        else:
            if ch.text:
                out.append(ch.text)
            walk(ch, out)
        if ch.tail:
            out.append(ch.tail)


def edition_text(root):
    out = []
    for d in root.iter(TEI + "div"):
        if d.get("type") == "edition":
            out.append(d.text or "")
            walk(d, out)
            out.append(" \n ")
    return "".join(out)


def title_date(t):
    m = re.search(r"(\d{3,4})-(\d\d)-(\d\d)", t)
    if m:
        return int(m.group(1))
    m = re.search(r"(\d{3,4})\s*(?:Ś|Saka|Śaka|Ś\.)", t)
    if m:
        return int(m.group(1)) + 78
    m = re.search(r"(\d{3,4})\s*CE", t)
    if m:
        return int(m.group(1))
    m = re.search(r"(\d{1,2})(?:st|nd|rd|th)(?:[-/](\d{1,2})(?:st|nd|rd|th))?\s*c(?:\.|entury)", t)
    if m:
        c1 = int(m.group(1))
        c2 = int(m.group(2)) if m.group(2) else c1
        return ((c1 - 1) * 100 + 1, c2 * 100)
    return None


def load_xlsx():
    recs = []
    base = os.path.join(DH, "IDENK-archive")
    wb = openpyxl.load_workbook(os.path.join(base, "DHARMA_IDENK_Metadata_Rest_V01.xlsx"), read_only=True)
    rows = list(wb["Work in progress"].iter_rows(values_only=True))
    hdr = rows[0]
    col = lambda sub: [i for i, h in enumerate(hdr) if h and sub in h][0]
    C = dict(name=col("Designation/Nama"), prov=col("Province"), kab=col("District/Township"),
             desa=col("Village/Neighborhood"), saka=col("Tahun Saka"), ce=col("Exact Date CE"))
    for r in rows[1:]:
        if r[1]:
            recs.append({k: r[i] for k, i in C.items()})
    wb = openpyxl.load_workbook(os.path.join(base, "DHARMA_IDENK_Metadata_Part1_20190919.xlsx"), read_only=True)
    for r in list(wb["Work in progress"].iter_rows(values_only=True))[4:]:
        if r[4]:
            recs.append(dict(name=r[4], prov=r[16], kab=r[15], desa=r[13], saka=r[66], ce=r[68]))
    idx = {}
    for r in recs:
        idx.setdefault(key(r["name"]), []).append(r)
    return recs, idx


def xlsx_date(h):
    try:
        if h["ce"] and str(h["ce"]).replace(".0", "").isdigit():
            return int(float(h["ce"])), "xlsx CE"
        if h["ce"]:
            m = re.search(r"(\d{3,4})\s*$", str(h["ce"]).replace("Sek. ", ""))
            if m:
                return int(m.group(1)), "xlsx CE"
        if h["saka"] and str(h["saka"]).replace(".0", "").isdigit():
            return int(float(h["saka"])) + 78, "xlsx Saka+78"
    except (ValueError, TypeError):
        pass
    return None, None


def region(prov, kab, desa, gnote):
    # Kabupaten only: matching on the desa string let "…Kulon" desa names in Wonogiri pass as Kulon Progo.
    s = str(kab or "").lower()
    if prov:
        if "timur" in str(prov).lower():
            return "East Java", "xlsx"
        if any(w in s for w in KEDU):
            return "Kedu", "xlsx"
        if any(w in s for w in PRAM) or "yogyakarta" in str(prov).lower():
            return "Prambanan-Mataram", "xlsx"
        if "tengah" in str(prov).lower():
            return "Other Central Java", "xlsx"
    g = gnote or ""
    if g.startswith("Kedu-Prambanan area"):
        return "Kedu/Prambanan (E082 area label)", "E082"
    if g in ("Kedu Plain", "Magelang, Kedu") or "Temanggung" in g:
        return "Kedu", "E082"
    if "Bantul" in g:
        return "Prambanan-Mataram", "E082"
    if "East Java" in g or "Dinoyo" in g:
        return "East Java", "E082"
    if g in ("Mataram Central Java", "Central Java"):
        return "Central Java (unspecific)", "E082-placeholder"
    if "Sindoro" in g:
        return "Other Central Java", "E082"
    return "Unlocated", ""


def in_window(d):
    if d is None:
        return False
    lo, hi = d if isinstance(d, tuple) else (d, d)
    return lo >= 701 and hi <= 1000


def main():
    recs, idx = load_xlsx()
    byname = {r["name"]: r for r in recs}
    geo = {r["filename"].replace("DHARMA_", "").replace(".xml", ""): r
           for r in csv.DictReader(open(E082, encoding="utf-8"))}
    insc, occ = [], []
    for f in sorted(glob.glob(os.path.join(DH, "xml", "*.xml"))):
        fid = os.path.basename(f)[7:-4]
        if fid.startswith("INSIDENKBorobudurHB"):
            continue
        root = etree.parse(f).getroot()
        title = re.sub(r"\s+", " ", "".join(root.find(".//" + TEI + "title").itertext()))
        d, src = title_date(title), "title"
        if d is None:
            src = None
            if fid in ALIAS:
                d, src = ALIAS[fid]
            else:
                hit = idx.get(key(re.sub(r"\(.*?\)", "", title))) or idx.get(key(fid.replace("INSIDENK", "")))
                if hit:
                    d, src = xlsx_date(hit[0])
        if not in_window(d):
            continue
        x = byname.get(XLSX_NAME.get(fid, "")) or (
            (idx.get(key(re.sub(r"\(.*?\)", "", title))) or idx.get(key(fid.replace("INSIDENK", ""))) or [{}])[0])
        g = geo.get(fid, {})
        reg, rsrc = region(x.get("prov"), x.get("kab"), x.get("desa"), g.get("geocode_note"))
        raw = TOK.findall(re.sub(r"\s+", " ", edition_text(root)))
        nt = [norm(t) for t in raw]
        n_v = 0
        for i, t in enumerate(nt):
            if not VILLAGE.match(t):
                continue
            n_v += 1
            j = i + 1
            if j < len(nt) and nt[j] in PART:
                j += 1
            prop = raw[j] if j < len(raw) else ""
            occ.append(dict(
                occ_id=f"{fid}#{n_v:03d}", inscription=fid, title=title,
                date=d if not isinstance(d, tuple) else f"{d[0]}-{d[1]}", date_src=src,
                region=reg, noun=raw[i], proposed_name=prop,
                kwic_left=" ".join(raw[max(0, i - 12):i]), kwic_noun=raw[i],
                kwic_right=" ".join(raw[i + 1:i + 13])))
        insc.append(dict(inscription=fid, title=title,
                         date=d if not isinstance(d, tuple) else f"{d[0]}-{d[1]}", date_src=src,
                         region=reg, region_src=rsrc, xlsx_kab=x.get("kab") or "", xlsx_desa=x.get("desa") or "",
                         n_village_nouns=n_v, n_tokens=len(raw)))
    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "inscriptions_701_1000.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(insc[0]))
        w.writeheader()
        w.writerows(insc)
    with open(os.path.join(OUT, "candidates_kwic.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(occ[0]))
        w.writeheader()
        w.writerows(occ)
    print("inscriptions in window:", len(insc), "| with village noun:", sum(1 for r in insc if r["n_village_nouns"]))
    print("title-dated:", sum(1 for r in insc if r["date_src"] == "title"))
    print("occurrences:", len(occ))
    print("by region (occurrences):", collections.Counter(o["region"] for o in occ).most_common())
    print("by region (inscriptions w/ noun):",
          collections.Counter(r["region"] for r in insc if r["n_village_nouns"]).most_common())


if __name__ == "__main__":
    main()
