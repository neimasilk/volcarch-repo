# Post-process the pandoc conversion of the SUBMITTED text so that citation wording matches what the
# revision uses ("and", not "&"; serial comma for three authors); content is not touched.
# Output = baseline.docx (no author, no personal properties).
# usage: python make_baseline.py RAW.docx OUT.docx [SUBMITTED.tex SUBMITTED.bib]   (with the last two: reference list in OL style)
import docx, re, sys
src, out = sys.argv[1], sys.argv[2]
d = docx.Document(src)
n = 0
SERIAL = re.compile(r"([A-Z][A-Za-z'’-]+, [A-Z][A-Za-z'’-]+) and ([A-Z])")

def fix_par(p):
    global n
    for r in p.runs:
        if " & " in r.text:
            r.text = r.text.replace(" & ", " and ")
            n += 1
        new = SERIAL.sub(lambda m: m.group(1) + ", and " + m.group(2), r.text)
        if new != r.text:
            r.text = new
            n += 1

for p in d.paragraphs:
    fix_par(p)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                fix_par(p)
# Reference list in the journal's style (since 2026-10-12 the revised article sets its list with ../ol_refs.py; the
# baseline uses the same formatter on the SUBMITTED .bib, so that the comparison shows which works were added, dropped or
# corrected and not a change of punctuation in every entry). The entries are those cited in the submitted .tex.
if len(sys.argv) > 4:
    sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent.parent))
    import ol_refs
    tex = open(sys.argv[3], encoding="utf-8").read()
    keys = []
    for m in re.finditer(r"\\(?:no)?cite[a-zA-Z]*\*?(?:\[[^\]]*\])*\{([^}]*)\}", tex):
        for k in m.group(1).split(","):
            k = k.strip()
            if k and k not in keys:
                keys.append(k)
    entries = ol_refs.reference_list(keys, bibs=[sys.argv[4]])
    bib = [p for p in d.paragraphs if p.style.name == "Bibliography"]
    if len(bib) != len(entries):
        raise SystemExit(f"reference list: {len(bib)} paragraphs in the conversion, {len(entries)} cited keys")
    for p, e in zip(bib, entries):
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        if p.runs:
            p.runs[0].text = e; p.runs[0].italic = None
        else:
            p.add_run(e)
        for el in list(p._p):                      # hyperlinks that pandoc made of the DOIs
            if el.tag.endswith("}hyperlink"):
                p._p.remove(el)
    print("reference list rewritten in the journal's style:", len(entries), "entries")

# Vertically merged cells of the submitted tables (LaTeX \multirow: "A (Full)", "B (Phon.)") are split into plain cells;
# the text stays in the first of them. Word's Compare records a changed vertical merge as a tracked cell merge, and
# accepting or rejecting such a change drops the text of the cells concerned: until 2026-10-08 that cost four words on
# "reject all"; on 2026-10-12 Word aligned that table with the revised Table 2 and "accept all" lost its six list names.
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
n_vm = 0
for el in list(d.element.body.iter(W + "vMerge")):
    el.getparent().remove(el)
    n_vm += 1
print("vertical merges removed from the baseline's tables:", n_vm)

cp = d.core_properties
cp.author = ""
cp.last_modified_by = ""
cp.title = "Submitted version (converted for comparison)"
cp.comments = ""
d.save(out)
print("runs changed:", n)
