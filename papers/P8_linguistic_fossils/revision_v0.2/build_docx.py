"""Build the P8 revision as a Word file in the Oceanic Linguistics template from the Markdown source.

usage: python build_docx.py [--anon]

Source: P8_revision_v0.4.md (this folder). Output: P8_revision_v0.4.docx, or P8_revision_v0.4_anonymous.docx with --anon
(author block and repository URL replaced by placeholders; acknowledgments left out, as the journal asks of an anonymized
file). Tables are filled from the result files of E228/E229/E231 at build time, so a number in a table can never drift from
its CSV. References are set by ol_refs.py in the journal's own style (Chicago 15th ed., author-date, as the template gives
it) from the two .bib files, for the keys listed in CITED:.

Markdown conventions (kept deliberately small):
  TITLE: / RUNNINGHEAD: / AUTHOR: name | affiliation | e-mail  (repeatable) / KEYWORDS: / CITED: key, key, ...
  ABSTRACT: followed by the abstract paragraph (ends at the first blank line)
  # A-level heading        ## B-level heading   (the OL styles number themselves; write no numbers)
  [[TABLE:T1|Caption]]     table from the result files; following lines beginning with † ‡ § or "Note." are table notes
                           (§ is the source's mark for the third note and is printed as #, the journal's third symbol)
  [[FIGURE:1|Caption]]     caption + position marker (the image file is uploaded separately, as the press asks)
  [[REFERENCES]]           the formatted reference list
  [[UNNUMBERED:Heading]]   an unnumbered heading (data availability, AI declaration, acknowledgments)
  *italic*  **bold**       inline; a reconstruction's asterisk is written \\* (e.g. \\*zalan) so it is not read as italics

Layout rules taken from the template (OL-template-1.dotx), checked 2026-10-12:
  - a heading that is followed by text runs on into it (Word "style separator": a hidden paragraph mark on the heading);
    a heading followed by another heading stands alone
  - tables: no grid; every other row shaded (D9D9D9); a rule under the last row; columns fitted to their contents; notes
    under the table in the style "OL footnote" with the marks † ‡ # after a tab; note marks superscript in the table
  - figures: caption, then the marker <figure n here>
"""
import csv, json, os, re, sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import ol_refs

REPO = HERE.parents[2]
E228 = REPO / "experiments/E228_p8_revision_analyses/results"
E229 = REPO / "experiments/E229_p8_revision_tables/results"
E231 = REPO / "experiments/E231_p8_what_coded_means/results"
ANON = "--anon" in sys.argv
# v0.4 (2026-10-12, final review): (1) the journal's house style (serial comma, numerals, lower-case cross-references,
# reference list in the journal's style, tables and headings as in the template, Figure 1 placed at its first mention);
# (2) corrections of content found by an independent adversarial read and its verification, each narrowing, qualifying
# or documenting a statement (list with reasons: PACKAGE_R1_20261012.md). No number of a table changes; Table 7 gains
# the distance of its last column. v0.3 stays in the folder as the file that sat in the portal from 2026-10-08.
VERSION = "v0.4"
SRC = HERE / f"P8_revision_{VERSION}.md"
OUT = HERE / (f"P8_revision_{VERSION}_anonymous.docx" if ANON else f"P8_revision_{VERSION}.docx")
if os.environ.get("P8_OUT"): OUT = Path(os.environ["P8_OUT"])   # build elsewhere when the target is open in Word
BASE = HERE / "ol_base.docx"          # the journal's .dotx with its content type changed to a document (see VENUE.md)

# ----------------------------------------------------------------------------- helpers
def rd(path):
    return list(csv.DictReader(open(path, encoding="utf-8")))

INLINE = re.compile(r"(\*\*.+?\*\*|\*(?!\s).+?(?<!\s)\*)")
MARKS = "†‡#"

def add_runs(p, text, size=None):
    text = text.replace("\\*", "\x00")           # protected asterisk (reconstructions)
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = p.add_run(piece[2:-2].replace("\x00", "*")); r.bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = p.add_run(piece[1:-1].replace("\x00", "*")); r.italic = True
        else:
            r = p.add_run(piece.replace("\x00", "*"))
        if size: r.font.size = Pt(size)
    return p

def para(d, text, style, **kw):
    p = d.add_paragraph(style=style)
    return add_runs(p, text, **kw)

def heading(d, text, style, run_in):
    """A heading; with run_in the paragraph mark is hidden (Word's style separator), so the next paragraph runs on."""
    p = para(d, text, style)
    if run_in:
        # the merged line takes its paragraph format from the heading, so the heading is justified like the text
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ppr = p._p.get_or_add_pPr()
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:vanish")); rpr.append(OxmlElement("w:specVanish"))
        ppr.append(rpr)
    return p

def cell_text(cell, text, style, d, bold=False):
    """Fill a table cell; a note mark at the end of the text (" †") is set superscript, as the template asks."""
    cell.text = ""
    p = cell.paragraphs[0]; p.style = d.styles[style]
    m = re.match(rf"^(.*\S)\s([{MARKS}])$", text)
    main, mark = (m.group(1), m.group(2)) if m else (text, "")
    if bold:
        r = p.add_run(main); r.bold = True
    else:
        add_runs(p, main)
    if mark:
        r = p.add_run(mark); r.font.superscript = True
    return p

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

def bottom_rule(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "auto")
    b.append(e)
    shd = tcpr.find(qn("w:shd"))
    if shd is not None: shd.addprevious(b)       # schema order: tcW, tcBorders, shd
    else: tcpr.append(b)

def auto_width(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    w = tcpr.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW"); tcpr.insert(0, w)
    w.set(qn("w:w"), "0"); w.set(qn("w:type"), "auto")

def table(d, caption, header, rows, notes=()):
    para(d, caption, "OL caption")
    t = d.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(header):
        cell_text(t.rows[0].cells[i], h, "OL table column heading", d, bold=True)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = cell_text(cells[i], str(v), "OL table contents", d)
            # a signed number must not wrap: Word breaks a line after the en-dash that serves as minus sign, and in the
            # narrow columns of table 8 the sign stood alone on one line and the number on the next (portal PDF, 12 Oct)
            if re.fullmatch(r"[+–−]\d[\d.,]*", str(v)):
                cells[i]._tc.get_or_add_tcPr().append(OxmlElement("w:noWrap"))
            if ri % 2 == 0: shade(cells[i], "D9D9D9")
            if ri == len(rows) - 1:
                bottom_rule(cells[i])
                # the style keeps table rows together with what follows; the last row releases it (as in the template)
                kn = OxmlElement("w:keepNext"); kn.set(qn("w:val"), "0"); p._p.get_or_add_pPr().append(kn)
    for row in t.rows:
        for c in row.cells:
            auto_width(c)
    para(d, "", "OL 3pt separator")
    for n in notes:
        if n[:1] == "§": n = "#" + n[1:]
        if n[:1] in MARKS:
            p = d.add_paragraph(style="OL footnote")
            add_runs(p, n[0] + "\t" + n[1:].strip())
        else:
            p = d.add_paragraph(style="OL footnote")
            p.paragraph_format.left_indent = Pt(0); p.paragraph_format.first_line_indent = Pt(0)
            add_runs(p, n)
    para(d, "", "OL 3pt separator")

def rnd(x, nd):
    q = Decimal(1).scaleb(-nd)
    return str(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))

def fmt3(x): return rnd(x, 3)

def fmtd(x, nd=3):
    v = rnd(x, nd)
    if float(v) == 0: return v.replace('-', '')
    return v.replace('-', '–') if v.startswith('-') else '+' + v   # the journal's minus sign is the en-dash (template)

def thou(x):
    return f"{int(x):,}"

# ----------------------------------------------------------------------------- tables from the result files
NAME = {"Muna": "Muna", "Bugis": "Bugis", "Makassar": "Makasar", "Wolio": "Wolio", "Toraja-Sadan": "Sa'dan Toraja", "Tolaki": "Tolaki"}
LABELS = {r["key"]: r["label"] for r in rd(E229.parent / "labels.csv")}

def T1():
    t1 = rd(E229 / "T1_label_by_list.csv")
    return (["List", "Forms", "Coded (n)", "Coded (%)", "Uncoded (n)", "Uncoded (%)"],
            [[r["list"], thou(r["forms"]), r["coded_n"], f"{float(r['coded_pct']):.1f}", r["candidate_n"], f"{float(r['candidate_pct']):.1f}"] for r in t1])

def T2():
    ci = {(r["list"], r["class"]): r for r in rd(E231 / "C_retention_intervals.csv")}
    def cell(lst, cls):
        r = ci[(lst, cls)]; return f"{rnd(100*int(r['n'])/int(r['base']), 1)} [{rnd(100*float(r['ci_lo']), 1)}, {rnd(100*float(r['ci_hi']), 1)}] ({r['n']})"
    rows = []
    for key in ("Makassar", "Bugis", "Toraja-Sadan", "Wolio", "Muna", "Tolaki"):
        rows.append([NAME[key], ci[(key, "uncoded")]["base"], cell(key, "retained_from_PMP"), cell(key, "coded_not_PMP"), cell(key, "uncoded")])
    return (["List", "Meanings with a PMP entry", "In the PMP entry's set, % [95% CI] (n)", "Coded in another set, % [95% CI] (n)", "No cognate set, % [95% CI] (n)"], rows)

def T3():
    t2 = rd(E229 / "T2_cv_performance.csv"); lab = {"FM25": "form + meaning (25)", "F17": "form only (17)"}
    rows = [[lab[r["input_set"]], r["classifier"], fmt3(r["auc_mean"]), fmt3(r["accuracy_mean"]), fmt3(r["accuracy_always_coded"]),
             fmt3(r["precision_candidate_mean"]), fmt3(r["recall_candidate_mean"]), fmt3(r["f1_candidate_mean"])] for r in t2]
    sd = "; ".join(f"{lab[r['input_set']]}, {r['classifier']}: {fmt3(r['auc_sd_seed_means'])} / {fmt3(r['auc_sd_folds'])}" for r in t2)
    return (["Inputs", "Classifier", "AUC", "Accuracy", "Always coded", "Precision", "Recall", "F1"], rows, sd)

def T4():
    t3 = rd(E229 / "T3_lolo.csv"); rows = []
    for r in [x for x in t3 if x["input_set"] == "FM25"]:
        f = next(x for x in t3 if x["input_set"] == "F17" and x["list_key"] == r["list_key"])
        if r["list_key"] == "SUMMARY":
            rows.append(["Mean (standard deviation) over the six lists", "", "", f"{fmt3(r['auc'])} ({fmt3(r['auc_sd_over_lists'])})", "", f"{fmt3(f['auc'])} ({fmt3(f['auc_sd_over_lists'])})", "", ""])
        else:
            rows.append([r["list"], r["n"], r["n_candidates"], fmt3(r["auc"]), fmt3(r["accuracy"]), fmt3(f["auc"]), fmt3(f["accuracy"]), fmt3(r["accuracy_majority_answer"])])
    return (["Held-out list", "n", "Uncoded", "AUC (F+M)", "Acc. (F+M)", "AUC (F)", "Acc. (F)", "Majority"], rows)

def T5():
    p3 = {r["contrast"]: r for r in rd(E231 / "P3_length_conditioned.csv") if r["strata"] == "list x number of letters"}
    eq = {"has_glottal": "written glottal mark", "sem_ACTION": "action meaning", "has_prefix_like": "onset string",
          "has_nasal_cluster": "nasal input as coded", "n_consonant_clusters": ">= 1 consonant-letter cluster"}
    rows = []
    for r in rd(E229 / "T6_profile.csv"):
        if r["type"] == "0/1":
            a, b = f"{100*float(r['mean_candidate']):.1f}%", f"{100*float(r['mean_coded']):.1f}%"; et = "odds ratio ‡"
        else:
            a, b = f"{float(r['mean_candidate']):.2f}", f"{float(r['mean_coded']):.2f}"; et = "mean difference †"
        eff = float(r["effect"]); sign = "+" if (r["type"] != "0/1" and eff > 0) else ""
        if r["property_key"] in eq:
            q = p3[eq[r["property_key"]]]; eqv = f"{rnd(q['odds_ratio'], 2)} [{rnd(q['ci_lo'], 2)}, {rnd(q['ci_hi'], 2)}]"
        else:
            eqv = "—"
        # the property names come from E229's labels.csv, the file that also labels Figure 1, so that the two cannot differ
        rows.append([LABELS.get(r["property_key"], r["property"]), a, b, r["n_lists_same_sign_as_pooled"], et, f"{sign}{rnd(eff, 2)} [{rnd(r['ci_lo'], 2)}, {rnd(r['ci_hi'], 2)}]", eqv])
    return (["Property", "Uncoded", "Coded", "Same sign (of 6)", "Effect type", "Effect [95% CI]", "At equal letter count #"], rows)

def T6():
    rows = [[r["list"], thou(r["n"]), r["candidate_and_profile"], r["candidate_and_no_profile"], r["coded_and_profile"], r["coded_and_no_profile"], f"{float(r['kappa']):.2f}"] for r in rd(E229 / "T5_cells_by_list.csv")]
    return (["List", "n", "Uncoded, profile", "Uncoded, no profile", "Coded, profile", "Coded, no profile", "Kappa"], rows)

def T7():
    rows, seen = [], {}
    for r in rd(E229 / "T7_examples_R1-11.csv"):
        if r["scope"] != "all six": continue
        k = r["cell"]; seen[k] = seen.get(k, 0)
        if seen[k] >= 2: continue
        seen[k] += 1
        # The nearest same-meaning form is printed with its distance: two of the eight are at 0.5, above the 0.34 of the
        # screen, so the column is not a column of look-alikes (until v0.3 it was headed "Nearest look-alike").
        look = (f"{r['sulawesi_lookalike_form']} ({r['sulawesi_lookalike_list']}; {float(r['sulawesi_lookalike_distance']):.2f})"
                if r["sulawesi_lookalike_form"] else "—")
        # the cells are named as in the heading of table 6 and in the text ("uncoded, profile")
        rows.append([k.replace("candidate", "uncoded").replace(" & ", ", "), NAME[r["language"]], r["concept"], r["form"] + (" ‡" if r["abvd_loan_flag"] == "1" else ""),
                     r["cognacy"] or "—", (r["pmp_forms"] or "—").replace("*", "\\*"), f"{float(r['p_candidate_oof25']):.2f}", look])
    return (["Cell", "List", "Meaning", "Form (ABVD)", "Set", "PMP (ABVD)", "Score", "Nearest same-meaning form (list; distance)"], rows)

def T8():
    names = {"V0_as_published": "As in the sources (ʔ or apostrophe)", "V1_q": "Mark written as q", "V2_k": "Mark written as k",
             "V3_unwritten": "Mark left unwritten", "V4_geminate": "Preglottalized consonant written as a geminate",
             "V5_all_as_glottal_letter": "Apostrophe written as ʔ (changes no input)", "V6_feature_removed": "Glottal input removed"}
    # Unrounded means (E229 09_table8_full_precision.py, checked there against E228 at four decimals). Until v0.2 the
    # held-out mean was taken from E228's per-list values, which are stored at three decimals, and two cells printed
    # -0.019 for an unrounded difference of -0.0197.
    g = rd(E229 / "T8_glottal_full_precision.csv")
    rows = [[names[r["convention"]], r["forms_changed"], fmt3(r["cv_auc_25"]), fmtd(r["delta_cv_25"]), fmt3(r["lolo_mean_25"]), fmtd(r["delta_lolo_25"])]
            for r in g]
    return (["Convention", "Forms changed", "AUC, cross-validated", "Δ", "AUC, held-out list (mean)", "Δ"], rows)

TABLES = {"T1": T1, "T2": T2, "T3": T3, "T4": T4, "T5": T5, "T6": T6, "T7": T7, "T8": T8}

# ----------------------------------------------------------------------------- build
def main():
    text = SRC.read_text(encoding="utf-8")
    meta = {"AUTHOR": []}
    body_lines = []
    abstract, in_abs = [], False
    for line in text.splitlines():
        m = re.match(r"^(TITLE|RUNNINGHEAD|AUTHOR|KEYWORDS|CITED|ABSTRACT):\s*(.*)$", line)
        if m and not body_lines:
            k, v = m.group(1), m.group(2).strip()
            if k == "AUTHOR": meta["AUTHOR"].append(v)
            elif k == "ABSTRACT": in_abs = True; abstract.append(v) if v else None
            else: meta[k] = v
            continue
        if in_abs:
            if line.strip() == "": in_abs = False
            else: abstract.append(line.strip())
            continue
        body_lines.append(line)
    keys = [k.strip() for k in re.split(r"[,\s]+", meta.get("CITED", "")) if k.strip()]

    d = Document(BASE)
    body = d.element.body
    for el in list(body):
        if not el.tag.endswith("}sectPr"):
            body.remove(el)

    # The template's odd-page header holds the words "Oceanic Linguistics Template" inside a table cell; replace that
    # text node with the running head (the even-page header carries the journal's volume line and stays).
    WT = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    for sec in d.sections:
        for hdr in (sec.header, sec.first_page_header, sec.even_page_header):
            try:
                for t in hdr._element.iter(WT):
                    if t.text and "template" in t.text.lower():
                        t.text = meta.get("RUNNINGHEAD", meta["TITLE"])
            except Exception:
                pass
    d.core_properties.title = meta["TITLE"]
    d.core_properties.author = "" if ANON else "; ".join(a.split("|")[0].strip() for a in meta["AUTHOR"])
    d.core_properties.last_modified_by = ""
    d.core_properties.comments = ""
    para(d, meta["TITLE"], "OL article title")
    if ANON:
        para(d, "[Authors anonymized for review]", "OL author name")
        para(d, "[Affiliations anonymized for review]", "OL author affiliation")
    else:
        for a in meta["AUTHOR"]:
            parts = [x.strip() for x in a.split("|")]
            para(d, parts[0], "OL author name")
            if len(parts) > 1: para(d, parts[1], "OL author affiliation")
    para(d, " ".join(abstract), "OL abstract")
    if meta.get("KEYWORDS"): para(d, "Keywords: " + meta["KEYWORDS"], "OL abstract")

    def next_is_text(j):
        """Is the next non-empty source line an ordinary paragraph (so that a heading may run on into it)?"""
        while j < len(body_lines) and body_lines[j].strip() == "":
            j += 1
        return j < len(body_lines) and not body_lines[j].lstrip().startswith(("#", "[["))

    first_after_heading = True
    after_run_in = False
    skipping = False                      # inside a section that the anonymized file leaves out
    i = 0
    paragraph_buf = []
    def flush():
        nonlocal paragraph_buf, first_after_heading, after_run_in
        if paragraph_buf:
            t = " ".join(paragraph_buf).strip()
            if ANON:
                t = re.sub(r"https://github\.com/[^\s,;)]+", "[repository URL redacted for review]", t)   # keeps the comma after it
                t = t.replace("https://doi.org/[DOI]", "[DOI redacted for review]")
                t = re.sub(r"https://doi\.org/10\.5281/zenodo\.\d+", "[data DOI redacted for review]", t)
            if not skipping:
                # after a run-in heading the text starts with a space on the heading's own line (template)
                para(d, (" " + t) if after_run_in else t, "OL text" if first_after_heading else "OL text indented")
            first_after_heading = False
            after_run_in = False
            paragraph_buf = []
    while i < len(body_lines):
        line = body_lines[i]
        s = line.strip()
        if s == "":
            flush(); i += 1; continue
        if s.startswith("## ") or s.startswith("# "):
            flush(); skipping = False
            level, txt = ("B", s[3:]) if s.startswith("## ") else ("A", s[2:])
            run_in = next_is_text(i + 1)
            heading(d, txt.rstrip(".") + ("" if txt.rstrip().endswith("?") else "."), f"OL heading {level}-level", run_in)
            first_after_heading = True; after_run_in = run_in; i += 1; continue
        m = re.match(r"^\[\[UNNUMBERED:(.+?)\]\]$", s)
        if m:
            flush()
            # "Please omit acknowledgements in the initial submission to maintain anonymity" (instructions for
            # contributors); in the final version they go into the first note.
            skipping = ANON and m.group(1).strip().lower().startswith("acknowledg")
            if not skipping:
                para(d, m.group(1).upper(), "OL references heading")
            first_after_heading = True; after_run_in = False; i += 1; continue
        m = re.match(r"^\[\[TABLE:(\w+)\|(.+?)\]\]$", s)
        if m:
            flush(); name, cap = m.group(1), m.group(2)
            notes = []
            j = i + 1
            while j < len(body_lines) and (body_lines[j].strip()[:1] in ("†", "‡", "§") or body_lines[j].strip().startswith("Note.")):
                notes.append(body_lines[j].strip()); j += 1
            spec = TABLES[name]()
            if name == "T3":
                header, rows, sd = spec; notes = notes + ["Standard deviation of the AUC over the 10 repetition means / over the 50 test sets: " + sd + "."]
            else:
                header, rows = spec
            if name == "T7" and any("‡" in str(c) for row in rows for c in row):
                notes = notes + ["‡ = flagged as a loan by ABVD."]
            table(d, cap, header, rows, notes)
            first_after_heading = False
            i = j; continue
        m = re.match(r"^\[\[FIGURE:(\d+)\|(.+?)\]\]$", s)
        if m:
            flush(); para(d, m.group(2), "OL caption")
            FIG = {"1": "F1_input_importance.png", "2": "F2_workflow.png", "3": "F3_pmp_classes.png", "4": "F4_heldout_auc.png", "5": "F5_glottal_by_list.png"}
            png = E229 / FIG.get(m.group(1), "none.png")
            if not ANON and png.exists():
                # reading copy only: the journal wants the figure as a separate file, which is uploaded beside the article
                pp = d.add_paragraph(style="OL text"); pp.add_run().add_picture(str(png), width=Pt(312))
            # the template's own form of the position marker
            para(d, f"<figure {m.group(1)} here>", "OL text")
            para(d, "", "OL 3pt separator"); first_after_heading = False; i += 1; continue
        if s == "[[REFERENCES]]":
            flush(); skipping = False; para(d, "REFERENCES", "OL references heading")
            for entry in ol_refs.reference_list(keys, markup=True):   # italics for journal and book titles, as in the template
                para(d, entry, "OL reference list")
            i += 1; continue
        paragraph_buf.append(s); i += 1
    flush()
    d.save(OUT)
    words = len(re.findall(r"\w+", " ".join(abstract) + " " + " ".join(l for l in body_lines if not l.startswith("[["))))
    print(f"saved {OUT.name}: paragraphs {len(d.paragraphs)}, tables {len(d.tables)}, abstract words {len(' '.join(abstract).split())}, body words about {words}, cited keys {len(keys)}")

if __name__ == "__main__":
    main()
